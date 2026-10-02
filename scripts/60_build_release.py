"""Assemble the public release as a clean folder, release/, that becomes the first commit of the public repository
(this working repository's history stays private: earlier commits hold school contact details).

  python 60_build_release.py            builds release/ (deleting any previous build)

Contents: the report as index.html (GitHub Pages), the pipeline scripts, the Phase 3 and 4 outputs, the small
openly licensed raw inputs, the public docs, and the files in release_template/ with their placeholders filled.
data/raw/MANIFEST.md is generated: every input the pipeline reads, with size, SHA-256, source and licence.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
import shutil
import zipfile

import importlib

import openpyxl

import config as C

merged = importlib.import_module("50_merged_report")

VERSION = "1.0.0"
DATE = dt.date.today().isoformat()
REPO = "JKeatingMU/commuting-to-college"
PAGES_URL = f"https://{REPO.split('/')[0].lower()}.github.io/{REPO.split('/')[1]}/"
DOI = None   # set once Zenodo mints it

DEST = C.ROOT / "release"
TEMPLATE = C.ROOT / "release_template"
SCRIPTS = ["config.py", "19_prepare_inputs.py", "20_origins_p3.py", "21_campuses_p3.py", "22_matrices_p3.py",
           "23_car_osrm_p3.py", "25_combine_p3.py", "26_accessibility_p3.py", "27_report_p3.py", "28_coach_routes_p3.py",
           "29_costs_p3.py", "30_cost_report.py", "31_schools_p3.py", "32_emissions_p3.py", "33_emissions_report.py",
           "34_burden_p3.py", "35_validation_p3.py", "40_hea_flows.py", "41_county_accessibility.py",
           "42_gravity_model.py", "43_p4_report.py", "44_cost_friction.py", "50_merged_report.py",
           "60_build_release.py"]
DOCS = ["DATA-SOURCES.md", "cost-model.md", "emissions-model.md", "environment.md", "requirements.lock.txt"]
STALE_OUTPUTS = {"p4_panel.csv"}
SCHOOLS_XLSX = "doe-data-on-individual-schools-postprimary-2025-26-20261002.xlsx"
SCHOOL_CONTACT_COLS = {"Principal Name", "Email", "Phone"}

# every raw input the published pipeline reads: (file, source, licence, shipped in the release)
RAW = [
    (f"cso-saps-2022-{C.SNAPSHOT}.zip", "CSO, Census 2022 Small Area Population Statistics, all geographies: "
     "https://www.cso.ie/en/census/census2022/census2022smallareapopulationstatistics/", "CC BY 4.0", False),
    (f"cso-small-areas-2022-{C.SNAPSHOT}.geojson", "Tailte Éireann / CSO, Small Areas 2022 (ungeneralised): "
     "https://data.gov.ie/dataset/cso-small-areas-national-statistical-boundaries-2022-ungeneralised1", "CC BY 4.0", False),
    (f"cso-builtup-areas-2022-{C.SNAPSHOT}.geojson", "Tailte Éireann / CSO, Built Up Areas 2022: https://data.gov.ie/",
     "CC BY 4.0", False),
    ("cso-f7145-2022-20261002.json", "CSO PxStat F7145 (JSON-stat 2.0): "
     "https://ws.cso.ie/public/api.restful/PxStat.Data.Cube_API.ReadDataset/F7145/JSON-stat/2.0/en", "CC BY 4.0", True),
    ("cso-f7065-2022-20261002.json", "CSO PxStat F7065 (JSON-stat 2.0): "
     "https://ws.cso.ie/public/api.restful/PxStat.Data.Cube_API.ReadDataset/F7065/JSON-stat/2.0/en", "CC BY 4.0", True),
    (f"tfi-gtfs-all-{C.SNAPSHOT}.zip", "National Transport Authority, Transport for Ireland GTFS (all operators): "
     "https://www.transportforireland.ie/transitData/Data/GTFS_All.zip (archived with the Zenodo record)", "CC BY 4.0", False),
    (f"ireland-osm-{C.SNAPSHOT}.osm.pbf", "OpenStreetMap via Geofabrik, ireland-and-northern-ireland extract of "
     f"{C.SNAPSHOT[:4]}-{C.SNAPSHOT[4:6]}-{C.SNAPSHOT[6:]}: https://download.geofabrik.de/europe/ireland-and-northern-ireland.html",
     "ODbL 1.0", False),
    ("ie-adm1-geoboundaries.geojson", "geoBoundaries IRL ADM1 (source Ordnance Survey Ireland): "
     "https://www.geoboundaries.org/api/current/gbOpen/IRL/ADM1/", "CC BY 4.0", True),
    ("ie-adm2-geoboundaries.geojson", "geoBoundaries IRL ADM2 (source Ordnance Survey Ireland): "
     "https://www.geoboundaries.org/api/current/gbOpen/IRL/ADM2/", "CC BY 4.0", True),
    ("pobal-hp-deprivation-ed-2022-20261002.csv", "Pobal HP Deprivation Index 2022, ED level: "
     "https://data.gov.ie/dataset/pobal-hp-deprivation-index-scores-2022", "CC BY 4.0", True),
    (SCHOOLS_XLSX, "Department of Education, Data on Individual Schools, post-primary, final 2025/26: "
     "https://www.gov.ie/en/department-of-education/collections/post-primary-schools-enrolment-figures/ "
     "(licence: https://data.gov.ie/dataset/data-on-individual-schools)", "CC BY 4.0", True),
    ("cao-computing-courses-2025-20260827.json", "CAO points charts 2025, computing (ISCED ICT) courses: "
     "https://www.cao.ie/", "No open licence", False),
] + [(f.name, "Higher Education Authority, Access our Data: Students, new entrants by county of domicile and "
      "institution, one ISCED field per file: https://hea.ie/statistics/data-for-download-and-visualisations/"
      "access-our-data/access-our-data-students/ (licence: "
      "https://data.gov.ie/dataset/hea-higher-education-student-and-graduate-data)", "CC BY 4.0", True)
     for f in sorted(C.RAW.glob("hea-*-newentrants-county-institute-*.txt"))]


def sha256(p) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def attribution() -> str:
    s = (C.ROOT / "docs" / "DATA-SOURCES.md").read_text()
    block = s[s.index("## 4."):s.index("## 5.")]
    return " ".join(l[1:].strip() for l in block.splitlines() if l.startswith(">"))


def citation_text() -> str:
    doi = f" https://doi.org/{DOI}" if DOI else " (DOI assigned on release through Zenodo)"
    return (f"Keating, J. G. (2026). *Commuting to College: Accessibility, Cost, Emissions and Choice for Higher "
            f"Education in Leinster* (Version {VERSION}) [Data set and report]. Maynooth University.{doi}. Individual research and analysis, not an "
            f"official Maynooth University publication.")


def strip_school_contacts(src, dst) -> None:
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        hdr = {c.value: c.column for c in ws[2] if c.value}
        for name in sorted(SCHOOL_CONTACT_COLS & set(hdr), key=lambda n: -hdr[n]):
            ws.delete_cols(hdr[name])
    fixed = dt.datetime(2026, 10, 2)
    wb.properties.created = wb.properties.modified = fixed
    wb.save(dst)
    # openpyxl stamps the current time into the zip entries; rewrite them with a fixed time so the
    # file, and its checksum in MANIFEST.md, are identical on every build
    with zipfile.ZipFile(dst) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            if name == "docProps/core.xml":   # openpyxl sets "modified" to now at save time
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-02T00:00:00Z", data)
            z.writestr(zipfile.ZipInfo(name, date_time=fixed.timetuple()[:6]), data, zipfile.ZIP_DEFLATED)
    left = {c.value for ws in openpyxl.load_workbook(dst).worksheets for c in ws[2]}
    assert not (SCHOOL_CONTACT_COLS & left), "contact columns survived"


def manifest() -> str:
    rows = ["# Raw inputs", "",
            "Every input the published pipeline reads, as downloaded. Files marked *in release* are in this folder;",
            "download the others from the source into `data/raw/` with the file name shown, and check the SHA-256.",
            f"The schools file here has the principal-name, email and phone columns removed (its checksum is of the",
            "released copy). This is individual research and analysis, not an official Maynooth University publication.", "",
            "| File | Size | SHA-256 | Source | Licence | In release |", "|---|---|---|---|---|---|"]
    for name, src, lic, ship in RAW:
        p = (DEST / "data" / "raw" / name) if ship else (C.RAW / name)
        if not p.exists():
            raise FileNotFoundError(p)
        rows.append(f"| `{name}` | {p.stat().st_size:,} | `{sha256(p)}` | {src} | {lic} | {'yes' if ship else 'no'} |")
    return "\n".join(rows) + "\n"


def fill(text: str) -> str:
    for k, v in {"VERSION": VERSION, "DATE": DATE, "REPO": REPO, "PAGES_URL": PAGES_URL,
                 "ATTRIBUTION": attribution(), "CITATION_TEXT": citation_text()}.items():
        text = text.replace("{" + k + "}", v)
    left = re.findall(r"\{[A-Z_]+\}", text)
    assert not left, f"unfilled placeholders: {left}"
    return text


def main() -> None:
    if DEST.exists():
        shutil.rmtree(DEST)
    for d in ("scripts/sections", "outputs", "data/raw", "docs"):
        (DEST / d).mkdir(parents=True)

    shutil.copyfile(C.OUT / "hei-commute.html", DEST / "index.html")
    for f in SCRIPTS:
        shutil.copyfile(C.ROOT / "scripts" / f, DEST / "scripts" / f)
    for f in (C.ROOT / "scripts" / "sections").glob("*.py"):
        shutil.copyfile(f, DEST / "scripts" / "sections" / f.name)
    for f in ("run_pipeline.sh", "r5py.yml"):
        shutil.copy2(C.ROOT / f, DEST / f)
    for f in DOCS:
        shutil.copyfile(C.ROOT / "docs" / f, DEST / "docs" / f)
    for f in sorted(C.OUT.iterdir()):
        if re.match(r"p[34]_.*\.(csv|json|png)$", f.name) and f.name not in STALE_OUTPUTS:
            shutil.copyfile(f, DEST / "outputs" / f.name)

    for name, _, _, ship in RAW:
        if ship and name == SCHOOLS_XLSX:
            strip_school_contacts(C.RAW / name, DEST / "data" / "raw" / name)
        elif ship:
            shutil.copyfile(C.RAW / name, DEST / "data" / "raw" / name)
    (DEST / "data" / "raw" / "MANIFEST.md").write_text(manifest())

    (DEST / "ABOUT.md").write_text(merged.about_markdown(merged.headlines()))
    for f in TEMPLATE.iterdir():
        (DEST / f.name).write_text(fill(f.read_text()))
    (DEST / ".gitignore").write_text(".venv/\n__pycache__/\n*.pyc\n.DS_Store\ndata/processed/\ndata/osrm/\n"
                                     + "".join(f"data/raw/{n}\n" for n, _, _, ship in RAW if not ship))
    (DEST / ".nojekyll").write_text("")

    files = [p for p in DEST.rglob("*") if p.is_file()]
    mb = sum(p.stat().st_size for p in files) / 1e6
    print(f"release/ built: {len(files)} files, {mb:.1f} MB, version {VERSION}, {PAGES_URL}")
    big = [(p.relative_to(DEST), round(p.stat().st_size / 1e6, 1)) for p in files if p.stat().st_size > 50e6]
    assert not big, f"files over GitHub's 50 MB warning: {big}"


if __name__ == "__main__":
    main()
